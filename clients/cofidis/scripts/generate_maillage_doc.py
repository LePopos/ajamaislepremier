from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Marges
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# Style de base
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(10)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14 if level == 1 else 11)
    run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e) if level == 1 else RGBColor(0x33, 0x33, 0x33)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    return p


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)


def add_table(doc, headers, rows, header_bg='1a1a2e', header_fg='FFFFFF'):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'

    # Header
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        set_cell_bg(cell, header_bg)
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)
        cell.paragraphs[0].paragraph_format.space_before = Pt(2)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)

    # Rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        bg = 'F5F5F5' if r_idx % 2 == 0 else 'FFFFFF'
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_bg(cell, bg)
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)

    doc.add_paragraph()
    return table


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(10)
        p.add_run(text).font.size = Pt(10)
    else:
        p.add_run(text).font.size = Pt(10)


# ── TITRE ──────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
run = p.add_run('Maillage interne — Règles par type de page')
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

add_body(doc, 'Sur chaque page de contenu, insérer des liens internes organisés en blocs. Ce document définit quels blocs mettre, combien de liens, et comment les sélectionner.')

doc.add_paragraph()

# ── 1. LES 4 BLOCS ────────────────────────────────────────────────────────
add_heading(doc, '1. Les 4 blocs disponibles')

add_table(doc,
    ['Bloc', 'Nom affiché (exemple)', 'Contenu'],
    [
        ['BLOC Q&A', '"Questions fréquentes sur le même sujet"', 'Pages Q&A du même thème'],
        ['BLOC GUIDE', '"Besoin d\'autres conseils sur le même thème ?"', 'Pages guides du même thème'],
        ['BLOC INCONTOURNABLES', '"À voir aussi"', '6 pages fixes (simulateurs, pages produit clés)'],
        ['BLOC OUTIL', '"Estimez votre projet"', 'Simulateur(s) du même thème'],
    ]
)

# ── 2. ACTIVATION PAR TYPE ────────────────────────────────────────────────
add_heading(doc, '2. Quels blocs activer selon le type de page')

add_table(doc,
    ['Type de page', 'BLOC Q&A', 'BLOC GUIDE', 'BLOC INCONTOURNABLES', 'BLOC OUTIL'],
    [
        ['Page Q&A',     '✅ 5 liens', '✅ 6 liens', '✅ 6 liens', '✅ 2 liens'],
        ['Page Guide',   '✅ 5 liens', '✅ 6 liens', '✅ 6 liens', '✅ 2 liens'],
        ['Page Produit', '✅ 5 liens', '✅ 6 liens', '✅ 6 liens', '✅ 2 liens'],
        ['Page Hub',     '—',          '✅ 6 liens', '✅ 6 liens', '✅ 2 liens'],
        ['Outil / Simulateur', '—',    '—',          '✅ 6 liens', '—'],
        ['Lexique',      '—',          '—',          '✅ 6 liens', '—'],
    ]
)

# ── 3. RÈGLES DE SÉLECTION ────────────────────────────────────────────────
add_heading(doc, '3. Comment sélectionner les liens de chaque bloc')

add_heading(doc, 'BLOC Q&A — 5 liens', level=2)
add_bullet(doc, 'Prendre les pages Q&A du même thème (ex : toutes les Q&A "crédit cuisine")')
add_bullet(doc, 'Exclure la page en cours')

add_heading(doc, 'BLOC GUIDE — 6 liens', level=2)
add_bullet(doc, 'Prendre les pages Guide du même thème')
add_bullet(doc, 'Si moins de 4 guides dans le thème : compléter avec des guides de thèmes proches (voir §4)')
add_bullet(doc, 'Exclure la page en cours')

add_heading(doc, 'BLOC INCONTOURNABLES — 6 liens fixes', level=2)
add_body(doc, 'Ces 6 pages apparaissent sur toutes les pages, sans exception :')
for url in [
    '/fr/credit/simulation-credit.html',
    '/fr/credit.html',
    '/fr/pret-personnel/pret-sur-mesure.html',
    '/fr/pret-personnel/simulation-pret.html',
    '/fr/credit/credit-renouvelable.html',
    '/fr/credit/credit-consommation.html',
]:
    add_bullet(doc, url)

add_heading(doc, 'BLOC OUTIL — 2 liens', level=2)
add_bullet(doc, 'Prendre le ou les simulateurs du même thème')
add_bullet(doc, 'Si aucun simulateur thématique : utiliser /fr/pret-personnel/simulation-pret.html par défaut')

doc.add_paragraph()

# ── 4. THÈMES ────────────────────────────────────────────────────────────
add_heading(doc, '4. Organisation par thème')

add_body(doc, 'Chaque page appartient à un thème, déterminé par son URL.')

add_table(doc,
    ['Thème', 'URLs concernées (début)', 'Thèmes proches (pour compléter BLOC GUIDE)'],
    [
        ['Rénovation énergétique', '/fr/pret-renovation-energetique/', 'Travaux, Crédit général'],
        ['Travaux',                '/fr/pret-personnel/pret-travaux/', 'Rénovation énergétique, Crédit cuisine, Crédit SDB'],
        ['Crédit cuisine',         '/fr/pret-personnel/credit-cuisine/', 'Travaux, Crédit déco'],
        ['Crédit auto',            '/fr/pret-personnel/credit-auto/ ou /fr/credit-auto/', 'Crédit moto'],
        ['Crédit moto',            '/fr/pret-personnel/credit-moto-scooter/', 'Crédit auto'],
        ['Crédit salle de bain',   '/fr/pret-personnel/credit-salle-de-bain/', 'Travaux, Crédit cuisine'],
        ['Crédit déco',            '/fr/pret-personnel/credit-decoration/', 'Crédit cuisine, Travaux'],
        ['Crédit piscine',         '/fr/pret-personnel/credit-piscine/', '—'],
        ['Crédit mariage',         '/fr/pret-personnel/credit-mariage/', '—'],
        ['Crédit voyage',          '/fr/pret-personnel/credit-voyage/', '—'],
        ['Prêt personnel (générique)', '/fr/pret-personnel/ (autres)', 'Crédit général'],
        ['Crédit général',         '/fr/credit/, /fr/guide-credit/', 'Prêt personnel, Rachat de crédit'],
        ['Rachat de crédit',       '/fr/rachat-de-credit/', 'Crédit général'],
    ]
)

# ── 5. LIVRABLE ───────────────────────────────────────────────────────────
add_heading(doc, '5. Livrable')

add_body(doc, 'Un fichier Excel cofidis_maillage_mapping.xlsx sera généré avec :')
add_bullet(doc, '', bold_prefix='Onglet Mapping : ')
# patch last bullet
doc.paragraphs[-1].runs[-1].text = 'pour chaque page, les URLs de chaque bloc prêtes à intégrer'
add_bullet(doc, '', bold_prefix='Onglet Inventaire : ')
doc.paragraphs[-1].runs[-1].text = 'toutes les URLs du site classées par type et par thème'

doc.save('/workspaces/ajamaislepremier/cahier_des_charges_maillage_cofidis.docx')
print('Done.')
